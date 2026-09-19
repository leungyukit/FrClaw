# -*- coding: utf-8 -*-
"""生成 Shopee 马来西亚站选品测算表（含自动公式/条件格式/下拉）"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ---------- 样式 ----------
HEAD_FILL = PatternFill("solid", fgColor="EE5A24")  # Shopee 橙
PARAM_FILL = PatternFill("solid", fgColor="FFF3E0")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")
WHITE_BOLD = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True)
THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ============================================================
# Sheet 2: 费率参数（先建，供公式引用）
# ============================================================
wp = wb.active
wp.title = "费率参数"
params = [
    ("参数", "数值", "说明（请按 Shopee 卖家学习中心最新规则修改）"),
    ("汇率（1 RM = ? RMB）", 1.65, "2025 年大致区间 1.60–1.72，下单结汇时以实时汇率为准"),
    ("佣金率", 0.04, "普通跨境卖家通常约 4%；新店免佣期可改为 0；商城/不同类目有差异"),
    ("交易手续费率", 0.0212, "约 2%（含 SST 后约 2.12%），以后台账单实际比例为准"),
    ("SLS 运费（每 10g，RMB）", 0.25, "简化估算值！实际按 SLS 马来价卡分区/分段计费，务必以后台价卡为准"),
    ("买家承担运费（RM）", 6.0, "即藏价/买家付运费部分，实际取决于重量段与包邮活动，可逐单在 I 列覆盖"),
    ("国内杂费+头程/单（RMB）", 1.0, "取货、打包、贴面单、送仓等分摊，已在 G 列逐行填写时此行仅作参考"),
    ("目标毛利率门槛", 0.30, "低于此值标红/提示利润不足"),
    ("竞品评价数门槛", 500, "前排商品评价数低于此值说明头部未垄断，更适合切入"),
    ("建议售价下限（RM）", 9, "低于此价扣完运费佣金后利润太薄，除非纯引流款"),
    ("建议售价上限（RM）", 60, "新手无评价，高于此价转化困难"),
]
for r, row in enumerate(params, 1):
    for c, v in enumerate(row, 1):
        cell = wp.cell(r, c, v)
        cell.border = THIN
        if r == 1:
            cell.font = WHITE_BOLD
            cell.fill = HEAD_FILL
            cell.alignment = CENTER
        else:
            cell.alignment = LEFT if c == 3 else CENTER
            if c == 2:
                cell.fill = PARAM_FILL
                cell.font = BOLD
wp["B3"].number_format = "0.00%"
wp["B4"].number_format = "0.00%"
wp["B8"].number_format = "0%"
wp.column_dimensions["A"].width = 26
wp.column_dimensions["B"].width = 12
wp.column_dimensions["C"].width = 70

# 参数行号映射（B 列）
RATE = "'费率参数'!$B$2"
COMM = "'费率参数'!$B$3"
FEE = "'费率参数'!$B$4"
PER10 = "'费率参数'!$B$5"
BUYER_SHIP = "'费率参数'!$B$6"
GM = "'费率参数'!$B$8"
REV_LIM = "'费率参数'!$B$9"
P_LO = "'费率参数'!$B$10"
P_HI = "'费率参数'!$B$11"

# ============================================================
# Sheet 1: 选品测算
# ============================================================
ws = wb.create_sheet("选品测算", 0)

headers = [
    ("序号", 6),
    ("品类", 14),
    ("商品名称/关键词\n(中/马来语)", 26),
    ("1688/货源链接", 20),
    ("拿货价\n(RMB)", 10),
    ("重量\n(g)", 8),
    ("国内杂费\n+头程(RMB)", 10),
    ("预估总运费\n(RMB)", 11),
    ("买家承担运费\n(RM，可改)", 12),
    ("卖家承担运费\n(RMB)", 11),
    ("售价\n(RM)", 9),
    ("平台佣金+\n手续费(RM)", 11),
    ("总成本\n(RMB)", 10),
    ("毛利\n(RM)", 9),
    ("毛利率", 9),
    ("月搜索量\n(参考)", 10),
    ("TOP20平均\n评价数", 10),
    ("TOP20均价\n(RM)", 10),
    ("价格优势\n(RM)", 9),
    ("测款建议", 14),
    ("状态", 11),
    ("备注", 22),
]
for c, (name, width) in enumerate(headers, 1):
    cell = ws.cell(1, c, name)
    cell.font = WHITE_BOLD
    cell.fill = HEAD_FILL
    cell.alignment = CENTER
    cell.border = THIN
    ws.column_dimensions[get_column_letter(c)].width = width
ws.row_dimensions[1].height = 34
ws.freeze_panes = "D2"

# 示例数据：品类, 名称, 链接, 拿货价, 重量g, 杂费, 售价RM, 搜索量, 评价数, TOP均价, 备注
samples = [
    ("穆斯林服饰", "Tudung Bawal 印花方巾", "（填1688链接）", 3.5, 35, 1.0, 12.9, 8000, 320, 15.9, "花色更新快，多色混批"),
    ("穆斯林服饰", "Telekung 女士祷告服（简约款）", "", 18.0, 350, 2.0, 39.9, 3500, 180, 49.0, "斋月前6-8周上架"),
    ("家居收纳", "桌面/冰箱折叠收纳盒", "", 8.0, 220, 1.5, 19.9, 5200, 460, 22.9, "主图展示收纳前后对比"),
    ("家居厨房", "厨房密封罐 3 件套", "", 6.0, 300, 1.5, 18.9, 2900, 610, 19.9, "注意防碎包装"),
    ("3C配件", "手机壳（小米/OPPO/vivo 热门机型）", "", 2.8, 60, 1.0, 9.9, 12000, 2200, 11.9, "头部评价多→红海产品，谨慎"),
    ("美妆工具", "美妆蛋+收纳架套装", "", 3.2, 40, 1.0, 12.9, 4400, 350, 14.9, "避开化妆品成品，只做工具"),
    ("母婴小件", "婴儿硅胶辅食餐具", "", 5.0, 150, 1.2, 16.9, 2100, 270, 21.5, "认准食品级硅胶资质"),
    ("女性配饰", "ins 风耳环/发饰多件套装", "", 2.5, 20, 0.8, 9.9, 6800, 410, 12.9, "轻小件运费优势大，适合凑单"),
]

START = 2
for i, s in enumerate(samples):
    r = START + i
    cat, name, link, cost, gram, misc, price, vol, reviews, topavg, note = s
    ws.cell(r, 1, i + 1)
    ws.cell(r, 2, cat)
    ws.cell(r, 3, name)
    ws.cell(r, 4, link)
    ws.cell(r, 5, cost)
    ws.cell(r, 6, gram)
    ws.cell(r, 7, misc)
    ws.cell(r, 8, f"=CEILING(F{r}/10,1)*{PER10}")
    ws.cell(r, 9, f"={BUYER_SHIP}")
    ws.cell(r, 10, f"=MAX(H{r}-I{r}*{RATE},0)")
    ws.cell(r, 11, price)
    ws.cell(r, 12, f"=K{r}*({COMM}+{FEE})")
    ws.cell(r, 13, f"=E{r}+G{r}+J{r}")
    ws.cell(r, 14, f"=K{r}-L{r}-M{r}/{RATE}")
    ws.cell(r, 15, f'=IF(K{r}=0,"",N{r}/K{r})')
    ws.cell(r, 16, vol)
    ws.cell(r, 17, reviews)
    ws.cell(r, 18, topavg)
    ws.cell(r, 19, f"=R{r}-K{r}")
    ws.cell(r, 20, (
        f'=IF(K{r}=0,"待填价",'
        f'IF(O{r}<{GM},"⚠️ 利润不足",'
        f'IF(OR(K{r}<{P_LO},K{r}>{P_HI}),"🔶 价格带不符",'
        f'IF(Q{r}>{REV_LIM},"🔴 竞争激烈",'
        f'"✅ 推荐测款"))))'
    ))
    ws.cell(r, 21, "待测评")
    ws.cell(r, 22, note)

# 空白行（公式预置到第 60 行）
MAX_ROW = 60
for r in range(START + len(samples), MAX_ROW + 1):
    ws.cell(r, 1, f'=IF(C{r}="","",ROW()-1)')
    ws.cell(r, 8, f'=IF(F{r}="","",CEILING(F{r}/10,1)*{PER10})')
    ws.cell(r, 9, f'=IF(F{r}="","",{BUYER_SHIP})')
    ws.cell(r, 10, f'=IF(F{r}="","",MAX(H{r}-I{r}*{RATE},0))')
    ws.cell(r, 12, f'=IF(K{r}="","",K{r}*({COMM}+{FEE}))')
    ws.cell(r, 13, f'=IF(E{r}="","",E{r}+N(G{r})+J{r})')
    ws.cell(r, 14, f'=IF(OR(K{r}="",E{r}=""),"",K{r}-L{r}-M{r}/{RATE})')
    ws.cell(r, 15, f'=IF(OR(K{r}="",K{r}=0),"",N{r}/K{r})')
    ws.cell(r, 19, f'=IF(OR(R{r}="",K{r}=""),"",R{r}-K{r})')
    ws.cell(r, 20, (
        f'=IF(C{r}="","",IF(K{r}="","待填价",'
        f'IF(O{r}<{GM},"⚠️ 利润不足",'
        f'IF(OR(K{r}<{P_LO},K{r}>{P_HI}),"🔶 价格带不符",'
        f'IF(N(Q{r})>{REV_LIM},"🔴 竞争激烈","✅ 推荐测款")))))'
    ))

# 格式
money_cols = [5, 7, 8, 10, 12, 13, 14, 18, 19]
for r in range(2, MAX_ROW + 1):
    for c in range(1, 23):
        cell = ws.cell(r, c)
        cell.border = THIN
        cell.alignment = CENTER if c != 3 and c != 4 and c != 22 else LEFT
    for c in money_cols:
        ws.cell(r, c).number_format = "0.00"
    ws.cell(r, 9).number_format = "0.0"
    ws.cell(r, 11).number_format = "0.00"
    ws.cell(r, 15).number_format = "0.0%"
    # 自动计算列底色
    for c in [1, 8, 10, 12, 13, 14, 15, 19, 20]:
        ws.cell(r, c).fill = CALC_FILL

# 下拉
dv_cat = DataValidation(type="list", formula1=(
    '"穆斯林服饰,女性配饰,家居收纳,家居厨房,3C配件,美妆工具,母婴小件,'
    '节日季节性,文具,汽车/摩托小件,宠物用品,其他"'), allow_blank=True)
dv_status = DataValidation(type="list",
    formula1='"待测评,测款中,加单补货,淘汰"', allow_blank=True)
ws.add_data_validation(dv_cat); ws.add_data_validation(dv_status)
dv_cat.add(f"B2:B{MAX_ROW}")
dv_status.add(f"U2:U{MAX_ROW}")

# 条件格式：毛利率
ws.conditional_formatting.add(f"O2:O{MAX_ROW}",
    CellIsRule(operator="greaterThanOrEqual", formula=["0.3"],
               fill=PatternFill("solid", fgColor="C6EFCE"),
               font=Font(color="006100", bold=True)))
ws.conditional_formatting.add(f"O2:O{MAX_ROW}",
    CellIsRule(operator="lessThan", formula=["0.15"],
               fill=PatternFill("solid", fgColor="FFC7CE"),
               font=Font(color="9C0006")))
# 建议列
ws.conditional_formatting.add(f"T2:T{MAX_ROW}",
    FormulaRule(formula=['ISNUMBER(SEARCH("推荐",T2))'],
                fill=PatternFill("solid", fgColor="C6EFCE"),
                font=Font(color="006100", bold=True)))
ws.conditional_formatting.add(f"T2:T{MAX_ROW}",
    FormulaRule(formula=['ISNUMBER(SEARCH("利润不足",T2))'],
                fill=PatternFill("solid", fgColor="FFEB9C")))
ws.conditional_formatting.add(f"T2:T{MAX_ROW}",
    FormulaRule(formula=['ISNUMBER(SEARCH("竞争激烈",T2))'],
                fill=PatternFill("solid", fgColor="FFC7CE")))

# ============================================================
# Sheet 3: 使用说明
# ============================================================
wg = wb.create_sheet("使用说明")
lines = [
    ("Shopee 马来西亚站 · 新手选品测算表 — 使用说明", True),
    ("", False),
    ("1. 你只需填白色列：B品类、C商品名、D链接、E拿货价、F重量(g)、G国内杂费、I买家承担运费（默认自动带出）、", False),
    ("   K售价(RM)、P月搜索量、Q TOP20平均评价数、R TOP20均价、U状态、V备注。灰色列为自动计算，勿手动改。", False),
    ("2. 所有费率集中在【费率参数】表：汇率、佣金率、交易手续费率、SLS 每 10g 运费、买家承担运费等，", False),
    ("   规则变动时只改参数表，全表自动重算。SLS 运费为简化估算，正式定价务必对照卖家后台最新价卡。", False),
    ("3. 核心公式：", False),
    ("   预估总运费 = CEILING(重量/10g) × 每10g费率", False),
    ("   卖家承担运费 = MAX(总运费 − 买家承担运费×汇率, 0)（即“藏价”成本）", False),
    ("   总成本(RMB) = 拿货价 + 国内杂费/头程 + 卖家承担运费", False),
    ("   毛利(RM) = 售价 − 售价×(佣金率+手续费率) − 总成本÷汇率", False),
    ("4. 测款建议判定逻辑：毛利率<30% → 利润不足；售价不在 RM9–60 → 价格带不符；", False),
    ("   TOP20 平均评价数>500 → 竞争激烈；全部通过 → ✅ 推荐测款。门槛均可在参数表调整。", False),
    ("5. 数据从哪来：P/Q/R 三列去 Shopee 马来前台搜关键词（tudung / sejadah / storage box 等），", False),
    ("   按销量排序看前 20 个商品；搜索量看卖家后台 Market Insight / 搜索框下拉，或知虾、ShopeeSpy。", False),
    ("6. 建议用法：首批铺 30–50 个 SKU、每款备货 5–10 件，借新品流量期+免运券测 1–2 周；", False),
    ("   状态列标记“测款中”，有自然出单且毛利率达标 → “加单补货”，否则“淘汰”，预算集中到 3–5 个潜力款。", False),
    ("7. 红线提醒：不碰侵权 IP/大牌同款、食品保健品、化妆品成品（马来需 NPRA 通报）、", False),
    ("   纯电池/充电宝/强磁/液体粉末（SLS 限运）、玻璃陶瓷大件、标准码鞋服（退货高）。", False),
    ("8. 节点提醒：斋月前 6–8 周开始上穆斯林节日款（tudung/telekung/家居装饰）；", False),
    ("   9.9/10.10/11.11/12.12 提前 2–3 周备货；年底雨季备雨伞雨衣防水鞋套。", False),
]
for i, (text, is_title) in enumerate(lines, 1):
    cell = wg.cell(i, 1, text)
    if is_title:
        cell.font = Font(bold=True, size=14, color="EE5A24")
    cell.alignment = Alignment(wrap_text=True, vertical="center")
wg.column_dimensions["A"].width = 110

out = "/Users/liangyj/workspace/frClaw/Shopee马来站_选品测算表.xlsx"
wb.save(out)
print("saved:", out)
